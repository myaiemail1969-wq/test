#!/usr/bin/env python3
# Generates a PowerPoint overview of Project Aeria. Run once, then delete or keep for future edits.

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Palette: dark slate background, amber accent (matches an "off-grid / signal fire" feel)
BG = RGBColor(0x12, 0x16, 0x1C)
PANEL = RGBColor(0x1C, 0x22, 0x2B)
ACCENT = RGBColor(0xE8, 0xA8, 0x4B)
TEXT = RGBColor(0xE8, 0xE6, 0xE1)
SUBTEXT = RGBColor(0xA9, 0xB1, 0xBA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return slide


def add_text(slide, left, top, width, height, text, size, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri", line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font
    return box


def add_bullets(slide, left, top, width, height, items, size=18, color=TEXT,
                accent_color=ACCENT, bullet="›"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run1 = p.add_run()
        run1.text = f"{bullet}  "
        run1.font.size = Pt(size)
        run1.font.color.rgb = accent_color
        run1.font.bold = True
        run1.font.name = "Calibri"
        run2 = p.add_run()
        run2.text = item
        run2.font.size = Pt(size)
        run2.font.color.rgb = color
        run2.font.name = "Calibri"
    return box


def accent_rule(slide, left, top, width, thickness=0.04):
    bar = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(thickness))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def kicker(slide, text):
    add_text(slide, 0.7, 0.45, 6, 0.4, text.upper(), 14, ACCENT, bold=True)


def panel(slide, left, top, width, height):
    box = slide.shapes.add_shape(1, Inches(left), Inches(top), Inches(width), Inches(height))
    box.fill.solid()
    box.fill.fore_color.rgb = PANEL
    box.line.color.rgb = RGBColor(0x2E, 0x36, 0x40)
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    return box


# ---------------------------------------------------------------------------
# 1. TITLE
# ---------------------------------------------------------------------------
s = add_slide()
add_text(s, 0.9, 2.5, 11.5, 0.5, "PROJECT AERIA", 20, ACCENT, bold=True)
add_text(s, 0.9, 3.0, 11.5, 1.6, "A Sovereign Offline Knowledge System", 44, TEXT, bold=True)
add_text(s, 0.9, 4.4, 10.5, 1.0,
         "A self-contained \"Civilization Restart Kit\": a local AI advisor, an encrypted personal\n"
         "vault, and a curated reference library — all running from a single drive, with zero\n"
         "internet connection required.",
         18, SUBTEXT, line_spacing=1.2)
accent_rule(s, 0.9, 5.85, 3.0)
add_text(s, 0.9, 6.05, 8, 0.4, "No cloud. No accounts. No signal required.", 15, SUBTEXT)

# ---------------------------------------------------------------------------
# 2. WHAT IS IT
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Overview")
add_text(s, 0.7, 0.85, 11, 0.9, "What Is Project Aeria?", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)
add_bullets(s, 0.7, 2.2, 6.1, 4.8, [
    "A USB drive that boots a complete offline AI system — no internet, no "
    "external services, no dependency on infrastructure that may not survive a crisis.",
    "Runs a local large language model (via llama.cpp) that answers questions "
    "using a private reference library stored on the same drive.",
    "Organized as a \"Council\" of domain-expert advisors — medicine, fabrication, "
    "and structural engineering — plus an encrypted personal vault for critical documents.",
    "Built to work cross-platform: Windows, macOS, and Linux launchers all boot "
    "the same system from the same drive.",
], size=16)
p = panel(s, 7.2, 2.2, 5.4, 4.8)
tf = p.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.35)
tf.margin_top = Inches(0.35)
add_text(s, 7.55, 2.55, 4.8, 0.4, "DESIGN PRINCIPLE", 14, ACCENT, bold=True)
add_text(s, 7.55, 3.05, 4.8, 3.7,
         "“If the grid goes down, the knowledge\n"
         "shouldn’t.”\n\n"
         "Every component — model, library, and\n"
         "vault — lives on the drive itself. Plug it\n"
         "into any compatible computer and the\n"
         "full system is available immediately,\n"
         "with nothing to install and nothing to\n"
         "phone home to.",
         16, SUBTEXT, line_spacing=1.25)

# ---------------------------------------------------------------------------
# 3. ARCHITECTURE / FOLDER LAYOUT
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Architecture")
add_text(s, 0.7, 0.85, 11, 0.9, "How the Drive Is Organized", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)

layout_rows = [
    ("00", "BOOT SYSTEM", "Cross-platform launchers (Windows / macOS / Linux), health checks, startup & shutdown scripts"),
    ("01", "THE BRAINS", "Local LLM model files (GGUF format) that power every advisor — runs entirely on-device"),
    ("02", "THE COUNCIL", "The advisor agents — Surgeon, Blacksmith, and Engineer — plus the chat interface that routes questions to them"),
    ("03", "THE ARCHIVES", "The reference library: field manuals, medical guides, engineering texts — indexed and searchable for instant retrieval"),
    ("04", "THE SCOUT", "Local vision tool — describes and analyzes images/photos entirely offline"),
    ("05", "USER ADDITIONS", "Space for the owner's own notes, documents, and custom material"),
    ("06", "PERSONAL VAULT", "AES-256 encrypted storage for family, medical, legal, and credential records"),
]
top = 2.15
row_h = 0.72
for i, (num, name, desc) in enumerate(layout_rows):
    y = top + i * row_h
    chip = slide_chip = s.shapes.add_shape(1, Inches(0.7), Inches(y), Inches(0.62), Inches(0.56))
    chip.fill.solid()
    chip.fill.fore_color.rgb = ACCENT
    chip.line.fill.background()
    chip.shadow.inherit = False
    tf = chip.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = num
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = BG
    add_text(s, 1.55, y + 0.03, 2.7, 0.5, name, 16, TEXT, bold=True)
    add_text(s, 4.35, y + 0.03, 8.4, 0.6, desc, 13.5, SUBTEXT, line_spacing=1.05)

# ---------------------------------------------------------------------------
# 4. THE COUNCIL
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "The Advisors")
add_text(s, 0.7, 0.85, 11, 0.9, "Meet the Council", 32, TEXT, bold=True)
add_text(s, 0.7, 1.6, 10.5, 0.5,
         "Three domain specialists, one local model — each with its own expert system prompt and knowledge focus.",
         16, SUBTEXT)

council = [
    ("THE SURGEON", "Medicine",
     "Trauma triage (MARCH protocol), wound management, fracture stabilization, "
     "improvised surgical care, pharmaceutical dosing, infection control, childbirth and "
     "long-term patient management."),
    ("THE BLACKSMITH", "Metallurgy & Fabrication",
     "Forge setup and heat treatment, metal identification (spark/file/magnet tests), "
     "fabrication from scrap, casting, forge and arc welding, and recognizing hazardous metals."),
    ("THE ENGINEER", "Structural & Civil Engineering",
     "Structural design and load calculations, water systems and filtration, off-grid "
     "power (solar, micro-hydro, generators), fortification and physical security, sanitation, "
     "and mechanical systems."),
]
card_w = 3.75
gap = 0.35
start_x = 0.7
for i, (title, domain, desc) in enumerate(council):
    x = start_x + i * (card_w + gap)
    card = panel(s, x, 2.4, card_w, 4.5)
    accent_rule(s, x + 0.3, 2.75, 1.0, thickness=0.06)
    add_text(s, x + 0.3, 2.95, card_w - 0.6, 0.5, title, 19, ACCENT, bold=True)
    add_text(s, x + 0.3, 3.45, card_w - 0.6, 0.4, domain.upper(), 12.5, SUBTEXT, bold=True)
    add_text(s, x + 0.3, 4.0, card_w - 0.6, 2.7, desc, 14, TEXT, line_spacing=1.25)

# ---------------------------------------------------------------------------
# 5. KEY FEATURES
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Capabilities")
add_text(s, 0.7, 0.85, 11, 0.9, "What It Can Do", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)

features = [
    ("100% Offline Operation", "No internet connection, cloud account, or external API calls — "
     "the model, the library, and the interface all run from the drive."),
    ("Grounded Answers from a Real Library", "Retrieval-augmented responses pull directly from "
     "an indexed archive of field manuals, medical references, and technical texts — not guesses."),
    ("Encrypted Personal Vault", "AES-256 encrypted storage for family, medical, legal, and "
     "credential documents, with no plaintext sensitive data ever stored on disk."),
    ("Scout — Offline Vision", "Send a photo to a local vision model for analysis (e.g. identifying "
     "a part, a plant, or an injury) without any image ever leaving the machine."),
    ("Cross-Platform Boot", "Identical experience on Windows, macOS, and Linux — plug in the "
     "drive and the matching launcher brings the whole system up."),
    ("Plug-and-Go Launch", "Root-level launcher scripts and drive autorun registration mean "
     "the system can start the moment the drive is connected."),
]
col_w = 5.85
row_h = 1.55
for i, (title, desc) in enumerate(features):
    col = i % 2
    row = i // 2
    x = 0.7 + col * (col_w + 0.4)
    y = 2.25 + row * row_h
    add_text(s, x, y, 0.4, 0.4, "◆", 16, ACCENT, bold=True)
    add_text(s, x + 0.4, y - 0.05, col_w - 0.4, 0.4, title, 17, TEXT, bold=True)
    add_text(s, x + 0.4, y + 0.42, col_w - 0.4, 1.0, desc, 13, SUBTEXT, line_spacing=1.15)

# ---------------------------------------------------------------------------
# 6. THE ARCHIVE LIBRARY
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Knowledge Base")
add_text(s, 0.7, 0.85, 11, 0.9, "The Archive Library", 32, TEXT, bold=True)
add_text(s, 0.7, 1.6, 11, 0.5,
         "Every answer the Council gives is grounded in real reference material stored and indexed on the drive.",
         16, SUBTEXT)

arch_cols = [
    ("MEDICAL", ["Where There Is No Doctor", "Where There Is No Dentist",
                 "A Book for Midwives", "Army First Aid (FM 4-25.11)",
                 "Anatomy & Physiology (OpenStax)"]),
    ("SURVIVAL & ENGINEERING", ["Army Survival Manual (FM 21-76)", "Engineer Field Data (FM 5-34)",
                                "Water Supply (FM 10-52)", "University & College Physics (OpenStax)",
                                "Chemistry: Atoms First (OpenStax)"]),
    ("HOW IT'S USED", ["PDFs and ZIM archives converted to indexed text",
                       "Full-text search across the whole library",
                       "Council agents cite and quote source material",
                       "Expandable — drop in any new reference and re-index"]),
]
for i, (head, items) in enumerate(arch_cols):
    x = 0.7 + i * 4.05
    add_text(s, x, 2.35, 3.8, 0.4, head, 14, ACCENT, bold=True)
    accent_rule(s, x, 2.8, 0.9, thickness=0.035)
    add_bullets(s, x, 3.05, 3.8, 4.0, items, size=13.5, bullet="–")

# ---------------------------------------------------------------------------
# 7. HARDWARE / WHAT YOU NEED
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Requirements")
add_text(s, 0.7, 0.85, 11, 0.9, "What It Runs On", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)

add_text(s, 0.7, 2.2, 5.6, 0.45, "THE DRIVE", 15, ACCENT, bold=True)
add_bullets(s, 0.7, 2.7, 5.6, 4.0, [
    "USB flash drive, 32 GB minimum (64 GB+ recommended)",
    "Formatted exFAT — readable on Windows, macOS, and Linux without drivers",
    "Holds the launcher, the model files, the archive library, and the vault",
], size=15)

add_text(s, 6.9, 2.2, 5.7, 0.45, "DRIVE SPACE PLANNING", 15, ACCENT, bold=True)
space_rows = [
    ("Runtime binaries (llama.cpp)", "~500 MB"),
    ("Fallback small model", "~1.5 GB"),
    ("Main model (7–8B, recommended)", "~4.7 GB"),
    ("Vision model + projector", "~1.8 GB"),
    ("Reference library", "remaining space"),
]
ry = 2.75
for label, size_ in space_rows:
    add_text(s, 6.9, ry, 4.3, 0.4, label, 14, TEXT)
    add_text(s, 11.2, ry, 1.4, 0.4, size_, 14, ACCENT, bold=True, align=PP_ALIGN.RIGHT)
    ry += 0.5
accent_rule(s, 6.9, ry + 0.05, 5.7, thickness=0.025)
add_text(s, 6.9, ry + 0.2, 4.3, 0.4, "Comfortable full kit", 14, TEXT, bold=True)
add_text(s, 11.2, ry + 0.2, 1.4, 0.4, "64 GB", 14, ACCENT, bold=True, align=PP_ALIGN.RIGHT)

add_text(s, 0.7, 5.1, 11.9, 0.4, "TARGET MACHINE", 15, ACCENT, bold=True)
add_text(s, 0.7, 5.6, 11.9, 1.2,
         "Any reasonably modern Windows, macOS, or Linux computer with a USB port. The target "
         "(\"survival\") machine never needs an internet connection — only the original setup "
         "machine does, to perform the one-time download.",
         15, SUBTEXT, line_spacing=1.25)

# ---------------------------------------------------------------------------
# 8. GETTING STARTED
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Setup")
add_text(s, 0.7, 0.85, 11, 0.9, "Getting Started", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)

steps = [
    ("1", "Format the drive", "Format a USB drive as exFAT so it works across Windows, macOS, and Linux."),
    ("2", "Download the components", "On a machine with internet: clone the project, download the model "
                                       "(Dolphin 2.9.1, an uncensored 8B model), and pull reference texts."),
    ("3", "Build the library", "Convert downloaded PDFs/ZIM archives to indexed text so the Council "
                                 "can search and cite them."),
    ("4", "Copy to the drive", "Place the boot system, model, library, and launcher scripts onto the drive."),
    ("5", "Plug in and launch", "Insert the drive into any target machine and run the launcher for that "
                                 "platform — Windows, macOS, or Linux — and the Council comes online."),
]
sy = 2.3
for num, title, desc in steps:
    chip = s.shapes.add_shape(9, Inches(0.7), Inches(sy), Inches(0.5), Inches(0.5))
    chip.fill.solid()
    chip.fill.fore_color.rgb = ACCENT
    chip.line.fill.background()
    chip.shadow.inherit = False
    tf = chip.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = num
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = BG
    add_text(s, 1.45, sy - 0.03, 3.0, 0.5, title, 16, TEXT, bold=True)
    add_text(s, 4.6, sy - 0.03, 8.0, 0.7, desc, 13.5, SUBTEXT, line_spacing=1.15)
    sy += 0.92

# ---------------------------------------------------------------------------
# 9. SECURITY & PRIVACY
# ---------------------------------------------------------------------------
s = add_slide()
kicker(s, "Trust")
add_text(s, 0.7, 0.85, 11, 0.9, "Security & Privacy by Design", 32, TEXT, bold=True)
accent_rule(s, 0.7, 1.75, 1.6)
add_bullets(s, 0.7, 2.3, 11.6, 4.5, [
    "No network calls anywhere in the runtime — the model, the search, and the vault all "
    "operate without ever reaching out to the internet.",
    "Personal documents are encrypted at rest with AES-256; nothing sensitive is ever "
    "written to disk in plaintext.",
    "No accounts, telemetry, or third-party services of any kind — the only copy of your "
    "data is the copy on your drive.",
    "Because everything is local and inspectable, you control exactly what the system "
    "knows and what it can access.",
], size=17)

# ---------------------------------------------------------------------------
# 10. CLOSING
# ---------------------------------------------------------------------------
s = add_slide()
add_text(s, 0.9, 2.7, 11.5, 1.3, "Plug it in. It's ready.", 40, TEXT, bold=True)
add_text(s, 0.9, 4.0, 10.3, 1.4,
         "Project Aeria packages an AI advisor, a reference library, and an encrypted vault "
         "onto a single drive — built to keep working when nothing else does.",
         18, SUBTEXT, line_spacing=1.3)
accent_rule(s, 0.9, 5.6, 3.0)
add_text(s, 0.9, 5.8, 8, 0.4, "PROJECT AERIA  —  Civilization Restart Kit", 15, ACCENT, bold=True)

out_path = "/home/user/test/Project_Aeria_Overview.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
